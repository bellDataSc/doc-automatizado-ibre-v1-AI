"""Gera o template .docx a partir do relatorio preenchido.

Decisoes aplicadas (confirmadas com a usuaria):
  1. Legendas de tabelas/graficos: MANTIDAS integralmente.
  2. Tabelas: cabecalhos mantidos + 1 unica linha de dados vazia.
  3. Rotulos fixos de linha: LIMPOS (so o cabecalho de coluna permanece).
  4. Campos variaveis da capa/cabecalhos: trocados por marcadores {{...}}.

Preserva: secoes e sectPr, orientacao e margens, cabecalhos/rodapes, logos,
estilos, numeracao automatica (campos SEQ), sumario e listas (campos TOC),
quebras de pagina e a sequencia de titulos/legendas/Fonte.
"""

import re
import shutil
import zipfile
from collections import Counter

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree

SRC = r"c:\Users\isabel.cruz\Desktop\DNIT_II.01.1_RR_15-08.2026_Tomo-1_Materiais-SICRO.docx"
DST = r"c:\Users\isabel.cruz\Desktop\Docs_Automatizado_IBRE_SICRO\materiais\template\TEMPLATE_DNIT_Tomo-1_Materiais-SICRO.docx"

# --- estilos de texto de conteudo: colapsar em 1 paragrafo vazio ---
CONTEUDO = {
    "20_Desenvolvimento-IBRE/SAPR",
    "Desenvolvimento",
    "21_Marcador 1-IBRE/SAPR",
    "37_Citacao-IBRE/SAPR",
    "07_RR-Corpo-Resumo-IBRE/SAPR",
}
# --- estilos que ancoram imagem de grafico: manter paragrafo, remover imagem ---
FIGURA = {"33_Figura-Grafico-IBRE/SAPR", "34_Fonte-Tab-Fig-Graf-IBRE/SAPR"}
# --- estilos da capa/folha de rosto: aplicar marcadores ---
CAPA = {"04_Folha-Rosto-IBRE/SAPR", "03_Detalhes-Capa-IBRE/SAPR"}

# Campos variaveis -> marcadores.
# GERAL: substituidos em qualquer posicao (capa, cabecalhos, rodapes).
MARC_GERAL = [
    (re.compile(r"Nº\s*\d{2}-\d{2}/\d{4}"), "Nº {{NUM_RELATORIO}}"),
    (re.compile(r"VDRI\s*/\s*\d+\s*/\s*\d{4}"), "{{PROTOCOLO_FGV}}"),
    (re.compile(r"REFERÊNCIA\s+[A-ZÀ-Ú]+/\d{4}"), "REFERÊNCIA {{REFERENCIA}}"),
    (re.compile(r"Referência\s+[A-ZÀ-Ú][a-zà-ú]+/\d{4}"), "Referência {{REFERENCIA}}"),
]
# EXATO: so quando o paragrafo inteiro e o campo (evita trocar anos no meio do texto).
MARC_EXATO = [
    (re.compile(r"^\s*[A-ZÀ-Ú][a-zà-ú]+ de \d{4}\s*$"), "{{MES_ANO_EMISSAO}}"),
    (re.compile(r"^\s*\d{4}\s*$"), "{{ANO}}"),
]

W_BR_PAGE = qn("w:br") + '[@' + qn("w:type") + '="page"]'
stats = Counter()


def tem_quebra_pagina(p):
    return bool(p.findall(".//" + W_BR_PAGE))


def eh_ancora_estrutural(p):
    """Paragrafo que NUNCA pode ser removido.

    O `w:sectPr` de cada secao mora no `w:pPr` do ultimo paragrafo da secao;
    remover esse paragrafo destroi a secao (orientacao, margens, cabecalho e
    rodape proprios). O mesmo vale para quebras de pagina.
    """
    ppr = p.find(qn("w:pPr"))
    if ppr is not None:
        if ppr.find(qn("w:sectPr")) is not None:
            return True
        if ppr.find(qn("w:pageBreakBefore")) is not None:
            return True
    return tem_quebra_pagina(p)


def limpar_runs(p, preservar_quebra=True):
    """Remove runs de texto/imagem do paragrafo, preservando pPr e quebras de pagina."""
    for filho in list(p):
        tag = filho.tag.split("}")[1]
        if tag == "pPr":
            continue
        if tag == "r":
            if preservar_quebra and filho.findall(W_BR_PAGE):
                for sub in list(filho):
                    st = sub.tag.split("}")[1]
                    if st == "rPr":
                        continue
                    if st == "br" and sub.get(qn("w:type")) == "page":
                        continue
                    filho.remove(sub)
                continue
            p.remove(filho)
        elif tag in ("hyperlink", "fldSimple", "smartTag", "ins", "del"):
            p.remove(filho)


def remover_imagens(p):
    n = 0
    for r in list(p.findall(qn("w:r"))):
        if r.findall(".//" + qn("w:drawing")) or r.findall(".//" + qn("w:pict")):
            p.remove(r)
            n += 1
    return n


def texto_paragrafo(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def definir_texto_xml(p, novo):
    """Reescreve o texto do paragrafo no nivel de `w:t`.

    Atua so nos elementos de texto: o primeiro recebe o conteudo novo e os
    demais ficam vazios. Nao mexe nos runs em si, de modo que `w:br`,
    `w:drawing` e a formatacao de caractere sobrevivem.
    """
    ts = list(p.iter(qn("w:t")))
    if not ts:
        return False
    ts[0].text = novo
    ts[0].set(qn("xml:space"), "preserve")
    for t in ts[1:]:
        t.text = ""
    return True


# ------------------------------------------------------------------ documento
shutil.copyfile(SRC, DST)
doc = Document(DST)
body = doc.element.body

# ---------- 1. paragrafos do corpo ----------
paras = body.findall(qn("w:p"))
estilo_de = {}
for par in doc.paragraphs:
    estilo_de[id(par._p)] = par.style.name

anterior_era_conteudo = False
for p in paras:
    est = estilo_de.get(id(p))

    if est in CONTEUDO:
        protegido = eh_ancora_estrutural(p)
        limpar_runs(p)
        primeiro_do_bloco = not anterior_era_conteudo
        if primeiro_do_bloco or protegido:
            stats["conteudo_esvaziado"] += 1
            if protegido:
                stats["conteudo_esvaziado_protegido"] += 1
            anterior_era_conteudo = True
        else:
            body.remove(p)
            stats["conteudo_removido"] += 1
        continue

    anterior_era_conteudo = False

    if est in FIGURA:
        k = remover_imagens(p)
        if k:
            stats["imagens_removidas"] += k
        continue

    if est in CAPA:
        txt = texto_paragrafo(p)
        novo = txt
        for rx, sub in MARC_GERAL:
            novo = rx.sub(sub, novo)
        for rx, sub in MARC_EXATO:
            if rx.match(novo):
                novo = sub
        if novo != txt and definir_texto_xml(p, novo):
            stats["capa_marcadores"] += 1

# ---------- 2. tabelas ----------
def linhas_de_cabecalho(t):
    """Cabecalho = linha 0; + linha 1 quando ha cabecalho de 2 niveis."""
    if len(t.rows) < 2:
        return 1
    try:
        r0 = [(c.text or "").strip() for c in t.rows[0].cells]
        r1 = [(c.text or "").strip() for c in t.rows[1].cells]
    except Exception:
        return 1
    # linha 0 totalmente mesclada = linha-titulo -> linha 1 tambem e cabecalho
    if r0 and len(set(r0)) == 1 and r0[0]:
        return 2
    # primeira coluna repetida = celula mesclada verticalmente no cabecalho
    if r0 and r1 and r0[0] and r0[0] == r1[0]:
        return 2
    return 1


def limpar_celula(cell):
    ps = cell._tc.findall(qn("w:p"))
    for i, p in enumerate(ps):
        limpar_runs(p, preservar_quebra=False)
        remover_imagens(p)
        if i > 0:
            cell._tc.remove(p)
    if not cell._tc.findall(qn("w:p")):
        cell._tc.append(cell._tc.makeelement(qn("w:p"), {}))


for t in doc.tables:
    n_hdr = linhas_de_cabecalho(t)
    linhas = list(t.rows)
    dados = linhas[n_hdr:]
    stats["tabelas"] += 1
    stats["linhas_cabecalho_mantidas"] += n_hdr

    if not dados:
        continue

    # primeira linha de dados: mantida e esvaziada
    vistos = set()
    for cell in dados[0].cells:
        if id(cell._tc) in vistos:
            continue
        vistos.add(id(cell._tc))
        limpar_celula(cell)
    stats["linhas_vazias_mantidas"] += 1

    # demais linhas: removidas
    for row in dados[1:]:
        tr = row._tr
        tr.getparent().remove(tr)
        stats["linhas_dados_removidas"] += 1

# ---------- 3. notas de rodape ----------
for rel in doc.part.rels.values():
    if "footnotes" not in rel.reltype:
        continue
    parte = rel.target_part
    raiz = parse_xml(parte.blob)
    for fn in raiz.findall(qn("w:footnote")):
        if int(fn.get(qn("w:id"), "0")) < 1:
            continue
        for p in fn.findall(qn("w:p")):
            limpar_runs(p, preservar_quebra=False)
        stats["notas_rodape_limpas"] += 1
    parte._blob = etree.tostring(
        raiz, xml_declaration=True, encoding="UTF-8", standalone=True
    )

# ---------- 4. cabecalhos e rodapes: marcadores ----------
# O texto do cabecalho/rodape do IBRE fica dentro de caixas de texto
# (w:txbxContent), fora do alcance de `hf.paragraphs` -> percorre o XML da parte.
CAMPO = (qn("w:fldChar"), qn("w:instrText"), qn("w:fldSimple"))

vistos_partes = set()
for sec in doc.sections:
    for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer,
               sec.even_page_header, sec.even_page_footer):
        parte = hf.part
        if id(parte) in vistos_partes:
            continue
        vistos_partes.add(id(parte))
        for p in parte.element.iter(qn("w:p")):
            # nunca tocar em paragrafo com campo (numero de pagina, PAGEREF...)
            if any(p.findall(".//" + tag) for tag in CAMPO):
                continue
            txt = texto_paragrafo(p)
            if not txt.strip():
                continue
            novo = txt
            for rx, sub in MARC_GERAL:
                novo = rx.sub(sub, novo)
            if novo != txt and definir_texto_xml(p, novo):
                stats["cabecalho_rodape_marcadores"] += 1

doc.save(DST)

# ---------- 5. descartar midia orfa ----------
# python-docx nao coleta as partes de imagem que deixaram de ser referenciadas:
# os 63 graficos continuariam no pacote como peso morto. Remove as relacoes de
# imagem sem uso e, em seguida, os arquivos de midia que nenhuma parte cita.
RELS_IMG = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def ids_referenciados(xml_bytes):
    raiz = etree.fromstring(xml_bytes)
    usados = set()
    for el in raiz.iter():
        for attr, val in el.attrib.items():
            if attr.startswith("{" + R_NS + "}") and val:
                usados.add(val)
    return usados


def limpar_midia(caminho):
    with zipfile.ZipFile(caminho) as z:
        partes = {n: z.read(n) for n in z.namelist()}

    # 1. por parte, descobrir quais r:id continuam em uso
    alvos_mantidos = set()
    rels_novos = {}
    for nome, dados in partes.items():
        if not nome.endswith(".rels") or nome == "_rels/.rels":
            continue
        dono = nome.replace("_rels/", "").replace(".rels", "")
        if dono not in partes:
            continue
        usados = ids_referenciados(partes[dono])
        raiz = etree.fromstring(dados)
        mudou = False
        for rel in list(raiz):
            if rel.get("Type") != RELS_IMG:
                continue
            if rel.get("Id") in usados:
                base = rel.get("Target").split("/")[-1]
                alvos_mantidos.add(base)
            else:
                raiz.remove(rel)
                mudou = True
        if mudou:
            rels_novos[nome] = etree.tostring(
                raiz, xml_declaration=True, encoding="UTF-8", standalone=True
            )

    # 2. reescrever o pacote sem as midias orfas
    descartadas = [
        n for n in partes
        if n.startswith("word/media/") and n.split("/")[-1] not in alvos_mantidos
    ]
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, dados in partes.items():
            if nome in descartadas:
                continue
            z.writestr(nome, rels_novos.get(nome, dados))
    return len(descartadas), len(alvos_mantidos)


orfas, mantidas = limpar_midia(DST)
stats["midias_orfas_descartadas"] = orfas
stats["midias_mantidas"] = mantidas

print("TEMPLATE GERADO:", DST)
for k in sorted(stats):
    print(f"  {k:<32} {stats[k]}")


