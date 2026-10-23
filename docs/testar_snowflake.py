"""Verificacao da origem Snowflake e da planilha das tabelas do RR, sem Snowflake.

Uma sessao falsa (SQLite em memoria) imita o que o app usa da sessao Snowpark:
sql(...).to_pandas(), table(...).schema.names, file.put_stream/get_stream e
write_pandas. Os dados sao SINTETICOS: as views tem colunas de exemplo
(REFERENCIA, TABELA, ...) so para exercitar a leitura; as colunas reais sao
as que a base devolver.

    python docs/gerar_fixture_graficos.py
    python docs/testar_snowflake.py
"""
import os
import re
import sqlite3
import sys
from io import BytesIO
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pandas as pd

import dnit_sicro_snowflake as sf
from materiais_readers import cruzar, ler_origem_precos, ler_resultado_final


class SessaoFalsa:
    def __init__(self):
        self.db = sqlite3.connect(":memory:", check_same_thread=False)
        self.db.create_function("STARTSWITH", 2, lambda a, b: int(str(a).startswith(b)))
        self.stage = {}
        self.file = SimpleNamespace(put_stream=self._put, get_stream=self._get)

    @staticmethod
    def _local(sql):
        return re.sub(r"BASES_SAPR\.AGENTE_SAPR\.", "", sql)

    def carregar(self, nome, df):
        df.to_sql(nome, self.db, index=False)

    def sql(self, query, params=None):
        q = self._local(query).strip()
        params = params or []

        def to_pandas():
            if q.startswith("LIST"):
                pasta = q.split("@", 1)[1].strip()
                nomes = [k for k in self.stage if k.startswith(pasta)]
                return pd.DataFrame({"name": nomes})
            if q.startswith("CALL"):
                return pd.DataFrame({"R": [f"carregada {params[0]}"]})
            return pd.read_sql_query(q, self.db, params=params)

        def collect():
            if q in ("BEGIN", "COMMIT", "ROLLBACK"):
                return []
            return to_pandas().to_dict("records")

        return SimpleNamespace(to_pandas=to_pandas, collect=collect)

    def table(self, nome):
        cur = self.db.execute(f"SELECT * FROM {self._local(nome)} LIMIT 0")
        return SimpleNamespace(schema=SimpleNamespace(names=[d[0] for d in cur.description]))

    def write_pandas(self, df, tabela, **_):
        df.to_sql(tabela, self.db, index=False, if_exists="append")

    def _put(self, fluxo, caminho, **_):
        self.stage[caminho.lstrip("@").replace("BASES_SAPR.AGENTE_SAPR.", "")] = fluxo.read()

    def _get(self, caminho, **_):
        return BytesIO(self.stage[caminho.lstrip("@").replace("BASES_SAPR.AGENTE_SAPR.", "")])


def faixas(tab, n, ref="2026-07"):
    q = [10 * (i + 1) for i in range(n)]
    return pd.DataFrame({"REFERENCIA": ref, "TABELA": tab, "ORDEM": range(1, n + 1),
                         "FAIXA": [f"f{i}" for i in range(n)], "QUANTIDADE": q,
                         "PROPORCAO": [x / sum(q) for x in q]})


def por_uf(tab, codigos):
    from materiais_config import UFS

    return pd.DataFrame([{"REFERENCIA": "2026-07", "REFERENCIA_ANTERIOR": "2026-04", "TABELA": tab,
                          "CODIGO": c, "ORDEM": o, "DESCRICAO": f"Item {c}", "UNIDADE": "kg", "UF": u,
                          "ORDEM_UF": ou, "PRECO_ANTERIOR": 10.0, "PRECO_ATUAL": 10.5,
                          "VARIACAO": 0.01 * ou * (1 if o % 2 else -1), "VAR_MAX": 0.27, "VAR_MIN": -0.27,
                          "VAR_MAX_ABS": 0.27, "RELACAO_LIDERANCA": "LÍDER", "COMERCIALIZACAO": "Ampla",
                          "ANALISTA": "Ana", "ENTRA_T12_T24": True}
                         for o, c in enumerate(codigos, 1) for ou, u in enumerate(UFS, 1)])


def resumo(tab, codigos):
    return pd.DataFrame([{"TABELA": tab, "REFERENCIA": "2026-07", "CODIGO": c, "DESCRICAO": f"Item {c}",
                          "UNIDADE": "m³", "COMERCIALIZACAO": "Ampla", "RELACAO_LIDERANCA": "LÍDER",
                          "ANALISTA": "Ana", "VAR_MAX": 0.21, "VAR_MIN": -0.05, "VAR_MAX_ABS": 0.21,
                          "MAX_ABS_UF": "SP", "UFS_ACIMA_10": 3, "QTD_UFS_ACIMA_20": 2, "UFS_ACIMA_20": "SP, RJ"}
                         for c in codigos])


T1920 = pd.DataFrame([{"REFERENCIA": "2026-07", "UFF_PC_PFM": 794, "UFF_PCT_PFM": 119,
                       "T19_PRECOS_UFF_DMT": 913, "T19_PRECOS_DMT_27UF": 24651, "T20_PCI": 157,
                       "T20_PRE": 131, "T20_PEI": 13, "T20_PRI": 18, "T20_TOTAL": 1232, "ITENS_RESTRITA": 1232}])
REFS_T26 = ["2025-07", "2025-10", "2026-01", "2026-04", "2026-07"]


def montar_sessao():
    s = SessaoFalsa()
    preco = ler_resultado_final("/tmp/graf_resultado_final.xlsx")
    origem = ler_origem_precos("/tmp/graf_origem_precos.xlsx")
    base = cruzar(preco, origem)
    longo = base.longo[["codigo", "descricao", "unidade", "uf", "preco", "origem"]]
    linhas = []
    for ref, fator in (("2026-01", 0.9), ("2026-04", 0.95), ("2026-07", 1.0)):
        d = longo.copy()
        d["preco"] = d["preco"] * fator
        d.insert(0, "REFERENCIA", ref)
        linhas.append(d)
    # um codigo de outra frente, que nao pode entrar em Materiais
    linhas.append(pd.DataFrame([{"REFERENCIA": "2026-07", "codigo": "E9001", "descricao": "x",
                                 "unidade": "h", "uf": "AC", "preco": 1.0, "origem": "PC"}]))
    precos = pd.concat(linhas).rename(columns=str.upper)
    s.carregar("SICRO_HIST_PRECO_UF", precos)
    s.carregar("SICRO_HIST_CONTROLE_REFERENCIA",
               pd.DataFrame({"REFERENCIA": ["2025-07", "2025-10", "2026-01", "2026-04", "2026-07"]}))
    s.carregar("SICRO_HIST_VW_TAB_FAIXAS", pd.concat(
        [faixas(t, n) for t, n in (("T10", 12), ("T13", 12), ("T14", 12), ("T17", 5), ("T21", 7))]
        + [faixas("T27", 3, r) for r in REFS_T26]))
    s.carregar("SICRO_HIST_VW_TAB_VARIACAO_ITEM", pd.concat(
        [por_uf("T11", ["M0003", "M0004"]), por_uf("T15", ["M0003"]),
         por_uf("T23", ["M0010"]), por_uf("T25", ["M0010", "M0011"])]))
    s.carregar("SICRO_HIST_VW_TAB_VAR_RESUMO", pd.concat(
        [resumo("T12", ["M0003", "M0004"]), resumo("T16", ["M0003"]),
         resumo("T24", ["M0010"]), resumo("T26", ["M0010", "M0011"])]))
    s.carregar("SICRO_HIST_VW_TAB_T19_T20", T1920)
    for v in ("SICRO_HIST_VW_SERIE_HIST", "SICRO_HIST_VW_AMPLITUDE", "SICRO_HIST_VW_VARIACAO"):
        s.carregar(v, pd.DataFrame({"REFERENCIA": ["2026-07"], "X": [1]}))
    s.carregar("SICRO_HIST_LISTA_RELEVANTE", pd.DataFrame({
        "REFERENCIA": ["2026-07"] * 2, "LISTA": ["T09", "T22"], "CODIGO": ["M0003", "M0004"],
        "ORDEM": [1, 1], "ARQUIVO_ORIGEM": ["x", "x"]}))
    s.carregar("SICRO_HIST_CADASTRO", pd.DataFrame({
        "REFERENCIA": ["2026-07"], "CODIGO": ["M0003"], "STATUS": ["Ativo"], "MODALIDADE": [None],
        "TIPO": ["Material"], "COMERCIALIZACAO": ["Ampla"], "RELACAO_LIDERANCA": [None],
        "CODIGO_LIDER": [None], "CLASSIFICACAO": [None], "FAMILIA_1": [None], "FAMILIA_2": [None],
        "ANALISTA": [None], "ARQUIVO_ORIGEM": ["x"]}))
    return s


def planilha_rr_sintetica() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Gerenciamento Atual"
    ws.append(["cadastro"])
    ws.append(["Status", "Externo(código)", "Elementar", "Descrição Externa", "modalidade",
               "Unidade", "Analista Atual", "Analista Ref-1", "Relação liderança", "Líder",
               "Tipo", "Comercialização", "classificacao", "familia_1", "familia_2"])
    ws.append(["Ativo", "M0003", "10001", "Aço", "x", "kg", "Ana", "Bia", "Líder", "M0003",
               "Material", "Ampla", "c", "Aço", "Aço CA"])
    ws.append(["Ativo", "M0004", "10002", "Areia", "x", "m3", "Ana", "Bia", "Liderado", "M0003",
               "Material", "Restrita", "c", "Areia", ""])
    ws.append([None] * 15)
    t = wb.create_sheet("Tabela 10")
    t.append(["Faixa", "Qtd"])
    t.append(["até 5%", 12])
    wb.create_sheet("Tabelas 19,20,21 e 22").append(["Metodologia", 794])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    s = montar_sessao()
    sf._sessao = s

    assert sf.configurado()
    assert sf.listar_referencias()[:3] == ["2026-07", "2026-04", "2026-01"]
    longo = sf.carregar_precos("2026-07")
    assert longo["codigo"].str.startswith("M").all(), "outra frente entrou em Materiais"
    print("── precos 2026-07:", longo["codigo"].nunique(), "insumos,", len(longo), "linhas")

    hist = sf.carregar_historico(["M0003"], "2026-07", 10)
    assert sorted(hist["referencia"].unique()) == ["2026-01", "2026-04", "2026-07"]   # so as que tem preco
    assert sorted(sf.carregar_historico(["M0003"], "2026-04", 10)["referencia"].unique()) == ["2026-01", "2026-04"]
    assert sf.carregar_historico(["M0003"], "2026-07", 2)["referencia"].nunique() == 2
    print("── historico ok")

    t10, separou = sf.recortar_tabela(sf.ler_view("tab_faixas", "2026-07"), "T10")
    assert separou and len(t10) == 12
    assert len(sf.ler_view_referencias("tab_faixas", REFS_T26[:2])) == 6
    t19, separou = sf.recortar_tabela(sf.ler_view("tab_t19_t20", "2026-07"), "T19")
    assert not separou and len(t19) == 1
    assert sf.numero_tabela("Tabela 10") == "T10" and sf.numero_tabela(27) == "T27"
    print("── views ok:", sf.estrutura_views()["view"].tolist())

    assert sf.ler_lista_relevante("T09", "2026-07")["CODIGO"].tolist() == ["M0003"]
    assert sf.referencias_com_cadastro() == ["2026-07"]

    # planilha das tabelas do RR
    from materiais_planilha_rr import abas_de_tabelas, abrir, ler_aba_bruta, ler_cadastro, resumo_cadastro

    livro = abrir(BytesIO(planilha_rr_sintetica()))
    cad = ler_cadastro(livro)
    assert cad.coluna_codigo == "Externo(código)", cad.coluna_codigo
    assert cad.completo, cad.avisos
    assert cad.dados["CODIGO"].tolist() == ["M0003", "M0004"]
    assert cad.dados["ANALISTA"].tolist() == ["Ana", "Ana"]
    assert cad.dados["CODIGO_LIDER"].tolist() == ["M0003", "M0003"]
    print("── cadastro:", resumo_cadastro(cad).to_dict("records"))
    abas = abas_de_tabelas(livro)
    assert abas == ["Tabela 10", "Tabelas 19,20,21 e 22"], abas
    assert ler_aba_bruta(livro, "Tabela 10").shape == (2, 2)
    livro.close()

    # carga controlada do cadastro: grava uma vez, recusa a segunda
    assert sf.inserir_cadastro("2026-10", cad.dados, "rr.xlsm") == 2
    try:
        sf.inserir_cadastro("2026-10", cad.dados, "rr.xlsm")
        raise AssertionError("gravou o cadastro duas vezes")
    except sf.ReferenciaJaExiste:
        pass
    assert sf.referencias_com_cadastro() == ["2026-10", "2026-07"]

    # stage + procedure
    sf.enviar_ao_stage("2026-10", "rr.xlsm", BytesIO(planilha_rr_sintetica()))
    assert sf.arquivos_no_stage("2026-10") == ["rr.xlsm"]
    assert abas_de_tabelas(abrir(BytesIO(sf.ler_do_stage(sf.caminho_stage("2026-10") + "rr.xlsm").read())))
    try:
        sf.carregar_referencia_pelo_stage("2026-07")
        raise AssertionError("aceitou referencia repetida")
    except sf.ReferenciaJaExiste:
        pass
    assert "2026-10" in sf.carregar_referencia_pelo_stage("2026-10")
    print("── carga ok")

    # grafico de evolucao com historico
    from materiais_charts import figura_evolucao

    base = cruzar(ler_resultado_final("/tmp/graf_resultado_final.xlsx"),
                  ler_origem_precos("/tmp/graf_origem_precos.xlsx"))
    base.historico = hist
    figura = figura_evolucao(base, "M0003", "julho/2026")
    rotulos = [t.get_text() for ax in figura.axes for t in ax.texts]
    assert "jan/2026" in rotulos and "jul/2026" in rotulos, rotulos
    legenda = [t.get_text() for t in figura.axes[0].get_legend().get_texts()]
    assert any("5,26%" in t for t in legenda), legenda   # 1.0 / 0.95 - 1
    print("── evolucao com 3 referencias ok")

    # tabelas prontas e escritas no Word (views com as colunas reais, dados sinteticos)
    import tempfile

    from docx import Document

    from materiais_tabelas import montar
    from materiais_writer import gerar_relatorio, linhas_da_tabela, tabelas_do_template

    da_base = {
        9: {"dados": pd.DataFrame({"CODIGO": [f"M{i:04d}" for i in range(1, 46)], "ORDEM": range(1, 46),
                                   "DESCRICAO": [f"Rel {i}" for i in range(1, 46)]}), "nota": ""},
        10: {"dados": faixas("T10", 12), "nota": ""}, 14: {"dados": faixas("T14", 12), "nota": ""},
        17: {"dados": faixas("T17", 5), "nota": ""},
        11: {"dados": por_uf("T11", ["M0003", "M0004"]), "nota": ""},
        15: {"dados": por_uf("T15", ["M0167"]), "nota": ""},
        12: {"dados": resumo("T12", ["M0043", "M9001"]), "nota": ""},
        25: {"dados": resumo("T26", [f"M{i:04d}" for i in range(50, 70)]), "nota": ""},
        19: {"dados": T1920, "nota": ""}, 20: {"dados": T1920, "nota": ""},
    }
    refs = REFS_T26
    prontas = montar(da_base, {12: {"M0043": "texto anterior"}},
                     pd.concat([faixas("T27", 3, r) for r in refs]), refs)
    assert sorted(prontas) == [9, 10, 11, 12, 14, 15, 17, 19, 20, 25, 26], sorted(prontas)
    assert prontas[19].celulas[(3, 0)] == ("(1) do total de 1.232 itens, 913 foram estabelecidos por "
                                           "frete metodológico e 319 por outras metodologias.")
    assert prontas[11].linhas[0][:4] == ["M0003", "Item M0003 – kg", "1,00%", "2,00%"]
    assert prontas[26].cabecalho == ["julho/25", "outubro/25", "janeiro/26", "abril/26", "julho/26"]

    with tempfile.TemporaryDirectory() as tmp:
        saida = os.path.join(tmp, "rel.docx")
        r = gerar_relatorio({}, saida, capa={"ANO": "2026"}, textos={}, tabelas=prontas)
        assert r.tabelas_preenchidas == sorted(prontas) and not r.tabelas_sem_slot
        t = tabelas_do_template(Document(saida))
    assert len(t[9]) == 2 and sum(len(p.rows) - 1 for p in t[9]) == 45       # cresceu
    assert len(t[12]) == 1 and len(t[12][0].rows) == 3                       # 6 partes -> 1
    assert len(t[25]) == 3 and sum(len(p.rows) - 1 for p in t[25]) == 20
    assert len(t[15]) == 1 and linhas_da_tabela(t[15][0])[1][2] == "1,00%"   # legenda depois da tabela
    t14 = linhas_da_tabela(t[14][0]) + linhas_da_tabela(t[14][1])
    assert t14[1][1] == "10" and t14[-2][1] == "780" and t14[-1][1] == "29"  # total e ÷27
    t17 = linhas_da_tabela(t[17][0])
    assert t17[5][1] == "100" and t17[6][1] == "50" and t17[7][1] == "150"   # total, zerados, total geral
    assert linhas_da_tabela(t[12][0])[1][4] == "texto anterior" and linhas_da_tabela(t[12][0])[2][4] == ""
    assert linhas_da_tabela(t[26][0])[1][1:] == prontas[26].cabecalho
    assert linhas_da_tabela(t[20][0])[6][1] == "1.232"
    print("── tabelas no Word ok:", r.tabelas_preenchidas)

    print("\nTODAS AS VERIFICACOES PASSARAM")
