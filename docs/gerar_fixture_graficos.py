"""Fixture voltado aos GRAFICOS: mesmo layout das planilhas reais, mas com os
codigos de insumo que o template cita nas legendas dos graficos, para que as
familias de grafico saiam de verdade na conferencia.

O fixture dos leitores (gerar_fixture.py) continua valendo e nao muda.
"""
import random
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook

from materiais_catalogo import catalogo
from materiais_config import UFS

random.seed(20260924)

RESULTADO_FINAL = "/tmp/graf_resultado_final.xlsx"
ORIGEM_PRECOS = "/tmp/graf_origem_precos.xlsx"

# Insumos citados nas legendas do template (evolucao e amplitude)
CODIGOS = {c for g in catalogo() for c in g.codigos}

# Familia da brita do Grafico 46, para conferir o ordenamento
FAMILIA_BRITA = ["M0005", "M0191", "M0192", "M0193", "M0808",
                 "M1097", "M1103", "M1135", "M2356"]
CODIGOS = sorted(CODIGOS | set(FAMILIA_BRITA))

UNIDADES = {}
ITENS = []
for indice, codigo in enumerate(CODIGOS):
    unidade = ["kg", "m³", "l", "un"][indice % 4]
    UNIDADES[codigo] = unidade
    ITENS.append((codigo, f"Insumo de teste {codigo}", unidade))
ITENS += [(f"E{9000 + i}", f"Equipamento de teste {i}", "un") for i in (1, 2, 3)]
ITENS += [("A9293", "Caminhao plataforma de teste", "un"),
          ("P9801", "Ajudante de teste", "h")]

ZERADOS = {"M0424", "M1790"}

# Fator por UF: faz os precos se moverem juntos entre insumos, como nos
# graficos reais (AC e RR caros, SP barato).
FATOR_UF = {uf: 1 + random.uniform(-0.12, 0.30) for uf in UFS}


def preco_de(codigo, uf):
    if codigo in ZERADOS:
        return 0.0
    base = 5 + (int(codigo[1:]) % 400)
    return round(base * FATOR_UF[uf] * random.uniform(0.985, 1.015), 4)


def origem_de(codigo, idx_uf):
    if codigo in ZERADOS:
        return "PRZ"
    if codigo.startswith("P"):
        return "SDP" if idx_uf % 3 else "SPC"
    return ["PFM", "PC", "PCI", "PMP", "PRE"][idx_uf % 5]


wb = Workbook()
ws = wb.active
ws.title = "Resultado Final"
ws["C3"] = "Resultado Final - Pesquisa Nacional de Precos - Ref.Julho de 2026"
for col, valor in enumerate(["CÓDIGO", "DESCRIÇÃO", "UNIDADE"] + UFS, start=1):
    ws.cell(row=4, column=col, value=valor)
linha = 5
for codigo, desc, uni in ITENS:
    ws.cell(row=linha, column=1, value=codigo)
    ws.cell(row=linha, column=2, value=desc)
    ws.cell(row=linha, column=3, value=uni)
    for j, uf in enumerate(UFS):
        ws.cell(row=linha, column=4 + j, value=preco_de(codigo, uf))
    linha += 1
wb.save(RESULTADO_FINAL)

wb = Workbook()
ws = wb.active
ws.title = "ENTREGA"
ws["D1"] = "Planilha de Origem Metodologica de Preco - Ref.Julho de 2026"
for col, valor in enumerate(
    ["CÓDIGO", "DESCRIÇÃO DO ITEM", "UNIDADE DO ITEM"] + UFS, start=1
):
    ws.cell(row=5, column=col, value=valor)
linha = 6
for codigo, desc, uni in ITENS:
    ws.cell(row=linha, column=1, value=codigo)
    ws.cell(row=linha, column=2, value=desc)
    ws.cell(row=linha, column=3, value=uni)
    for j, _uf in enumerate(UFS):
        ws.cell(row=linha, column=4 + j, value=origem_de(codigo, j))
    linha += 1
oculta = wb.create_sheet("Tabela RGB utilizada")
oculta["A1"] = "nomenclatura antiga - nao usar"
oculta.sheet_state = "hidden"
wb.save(ORIGEM_PRECOS)

print(f"{len(ITENS)} insumos, {len(CODIGOS)} do catalogo do template")
print(RESULTADO_FINAL)
print(ORIGEM_PRECOS)
