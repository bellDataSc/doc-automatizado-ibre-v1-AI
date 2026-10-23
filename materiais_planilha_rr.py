"""Leitura da planilha das tabelas do RR (SICRO_AAAA.MM_TABELAS-MATERIAIS-RR-TOMO-1).

E o .xlsm pesado (centenas de MB descompactado, ~1,9 milhao de formulas) em
que o analista monta as tabelas do relatorio. Aqui ele e aberto em modo
somente leitura (`read_only=True`) e com os valores ja calculados pelo Excel
(`data_only=True`): nenhuma formula e recalculada e o VBA nao e executado.
Sao lidas so as abas pedidas, linha a linha, sem carregar o arquivo inteiro
na memoria em forma de planilha.

Dois usos:
  - cadastro (aba "Gerenciamento Atual") para alimentar SICRO_HIST_CADASTRO,
    que e o que falta na base para as tabelas de outras referencias;
  - as abas "Tabela ..." como o analista as deixou, para conferencia ou
    para entrada manual quando a base nao tiver a referencia.

Nenhuma regra de negocio aqui: le, normaliza cabecalho e devolve DataFrames.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field

import pandas as pd

from materiais_config import PADRAO_CODIGO, PLANILHA_RR

_CODIGO = re.compile(PADRAO_CODIGO)


def _norm(texto) -> str:
    """Cabecalho comparavel: sem acento, minusculo, espacos simples."""
    if texto is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().lower()


def _texto(valor) -> str:
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()


def abrir(fonte):
    """Abre o .xlsm/.xlsx sem recalcular nada. `fonte` e caminho ou arquivo."""
    from openpyxl import load_workbook

    return load_workbook(fonte, read_only=True, data_only=True, keep_links=False)


def abas(livro) -> list[str]:
    return list(livro.sheetnames)


# ──────────────────────────────────────────────────────────────
# CADASTRO
# ──────────────────────────────────────────────────────────────
@dataclass
class Cadastro:
    dados: pd.DataFrame                  # colunas de SICRO_HIST_CADASTRO (sem REFERENCIA)
    coluna_codigo: str                   # cabecalho usado como CODIGO
    cabecalhos: dict                     # coluna de destino -> cabecalho lido
    avisos: list[str] = field(default_factory=list)

    @property
    def completo(self) -> bool:
        return not any(a.startswith("Coluna") for a in self.avisos)


def ler_cadastro(livro) -> Cadastro:
    aba = PLANILHA_RR["aba_cadastro"]
    if aba not in livro.sheetnames:
        raise ValueError(f"A planilha não tem a aba '{aba}'. Abas: {livro.sheetnames[:8]}...")

    linha_cab = PLANILHA_RR["linha_cabecalho_cadastro"]
    linhas = livro[aba].iter_rows(min_row=linha_cab, values_only=True)
    cabecalho = [_texto(v) for v in next(linhas)]
    corpo = [list(l) for l in linhas if any(v is not None for v in l)]
    df = pd.DataFrame(corpo, columns=[c or f"_col{i}" for i, c in enumerate(cabecalho)])
    normalizados = {_norm(c): c for c in df.columns}

    # CODIGO: entre as candidatas, a que mais traz codigos no padrao SICRO
    melhor, acertos = "", -1
    for candidata in PLANILHA_RR["candidatas_codigo"]:
        original = normalizados.get(candidata)
        if original is None:
            continue
        n = int(df[original].map(lambda v: bool(_CODIGO.match(_texto(v)))).sum())
        if n > acertos:
            melhor, acertos = original, n
    if not melhor or acertos <= 0:
        raise ValueError(
            "Não achei a coluna de código (Externo/Elementar) com códigos no padrão "
            f"SICRO na aba '{aba}'. Cabeçalho lido: {cabecalho[:10]}..."
        )

    avisos, usados = [], {}
    saida = pd.DataFrame({"CODIGO": df[melhor].map(_texto)})
    for destino, aceitos in PLANILHA_RR["colunas_cadastro"].items():
        original = next((normalizados[a] for a in aceitos if a in normalizados), None)
        if original is None:
            avisos.append(f"Coluna '{destino}' não encontrada (procurei {aceitos}).")
            saida[destino] = None
            continue
        usados[destino] = original
        saida[destino] = df[original].map(lambda v: _texto(v) or None)

    validos = saida["CODIGO"].map(lambda c: bool(_CODIGO.match(c)))
    descartadas = int((~validos).sum())
    saida = saida[validos].reset_index(drop=True)
    if descartadas:
        avisos.append(f"{descartadas} linha(s) sem código válido ignoradas.")
    duplicados = int(saida["CODIGO"].duplicated().sum())
    if duplicados:
        avisos.append(f"Coluna CODIGO: {duplicados} código(s) repetido(s) no cadastro.")

    return Cadastro(dados=saida, coluna_codigo=melhor, cabecalhos=usados, avisos=avisos)


def resumo_cadastro(cadastro: Cadastro) -> pd.DataFrame:
    """Contagem por TIPO x COMERCIALIZACAO, para conferir com a base."""
    df = cadastro.dados.fillna({"TIPO": "(vazio)", "COMERCIALIZACAO": "(vazio)"})
    return (
        df.groupby(["TIPO", "COMERCIALIZACAO"]).size()
        .reset_index(name="codigos").sort_values(["TIPO", "COMERCIALIZACAO"])
        .reset_index(drop=True)
    )


# ──────────────────────────────────────────────────────────────
# ABAS DAS TABELAS, COMO O ANALISTA DEIXOU
# ──────────────────────────────────────────────────────────────
def abas_de_tabelas(livro) -> list[str]:
    prefixo = _norm(PLANILHA_RR["prefixo_abas_tabelas"])
    return [a for a in livro.sheetnames if _norm(a).startswith(prefixo)]


def ler_aba_bruta(livro, aba: str, max_linhas: int = 400, max_colunas: int = 60) -> pd.DataFrame:
    """Valores calculados de uma aba, sem interpretar cabecalho.

    Linhas e colunas inteiramente vazias saem. Os rotulos de linha/coluna sao
    os do Excel (1, 2, ... / A, B, ...) para o analista achar a celula.
    """
    from openpyxl.utils import get_column_letter

    linhas = []
    for indice, linha in enumerate(
        livro[aba].iter_rows(max_row=max_linhas, max_col=max_colunas, values_only=True),
        start=1,
    ):
        linhas.append((indice, list(linha)))
    if not linhas:
        return pd.DataFrame()

    largura = max(len(l) for _, l in linhas)
    df = pd.DataFrame(
        [l + [None] * (largura - len(l)) for _, l in linhas],
        index=[i for i, _ in linhas],
        columns=[get_column_letter(c + 1) for c in range(largura)],
    )
    df = df.dropna(how="all").dropna(axis=1, how="all")
    return df.astype(object).where(df.notna(), "")
