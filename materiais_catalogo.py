"""Catalogo dos graficos do Tomo 1 - Materiais, lido do proprio template Word.

O template TEMPLATE_DNIT_Tomo-1_Materiais-SICRO.docx traz, na ordem do
documento, a legenda de cada grafico:

    Grafico 1  - Evolucao dos precos do Item M0003 Aco CA 25
    Grafico 42 - Ordenamento da familia Aco CA
    Grafico 51 - Precos e amplitude Item M0427 Cordoalha nua ...

Este modulo le essas legendas e devolve a lista dos 63 graficos com o tipo e,
quando a legenda traz, o codigo do insumo. Nada e inventado aqui: o catalogo e
exatamente o que esta escrito no template. Trocar o template troca o catalogo.

Os graficos de ordenamento sao por FAMILIA, e a legenda nao diz quais insumos
compoem cada familia - esse dado nao esta nas duas planilhas de entrada nem no
template. Eles saem com `codigos` vazio e `pendente` preenchido.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from functools import lru_cache

TEMPLATE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "template",
    "TEMPLATE_DNIT_Tomo-1_Materiais-SICRO.docx",
)

# Estilos do template (FGV IBRE)
ESTILO_LEGENDA = ("28_Titulo-Tab-Fig-Graf-IBRE/SAPR", "Caption")
ESTILO_FIGURA = "33_Figura-Grafico-IBRE/SAPR"
ESTILO_LISTA = "table of figures"   # indice; nao e o corpo do documento

EVOLUCAO = "evolucao"
AMPLITUDE = "amplitude"
ORDENAMENTO = "ordenamento"
OUTRO = "outro"

_RE_LEGENDA = re.compile(r"^Gráfico\s+(\d+)\s*[-–]\s*(.*)$")
_RE_EVOLUCAO = re.compile(r"^Evolução dos preços do Item\s+([AEMP]\d+)\s+(.*)$")
_RE_AMPLITUDE = re.compile(r"^Preços e amplitude Item\s+([AEMP]\d+)\s+(.*)$")
_RE_ORDENAMENTO = re.compile(r"^Ordenamento\s+(.*)$")


@dataclass
class Grafico:
    """Um grafico do relatorio, como o template o descreve."""

    numero: int
    tipo: str
    legenda: str                      # texto completo, sem o "Grafico N - "
    codigos: list = field(default_factory=list)
    pendente: str = ""                # o que falta para gerar, quando falta

    @property
    def titulo(self) -> str:
        return f"Gráfico {self.numero} - {self.legenda}"

    @property
    def geravel(self) -> bool:
        return not self.pendente


def _classificar(numero: int, texto: str) -> Grafico:
    m = _RE_EVOLUCAO.match(texto)
    if m:
        return Grafico(numero, EVOLUCAO, texto, [m.group(1)])

    m = _RE_AMPLITUDE.match(texto)
    if m:
        return Grafico(numero, AMPLITUDE, texto, [m.group(1)])

    if _RE_ORDENAMENTO.match(texto):
        return Grafico(
            numero,
            ORDENAMENTO,
            texto,
            [],
            pendente="composição da família e relação líder/liderado",
        )

    return Grafico(numero, OUTRO, texto, [], pendente="base fora das duas planilhas")


@lru_cache(maxsize=1)
def catalogo(caminho_template: str = TEMPLATE) -> tuple:
    """Le o template e devolve os graficos na ordem do documento."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    doc = Document(caminho_template)
    vistos: dict = {}

    for filho in doc.element.body.iterchildren():
        if filho.tag != qn("w:p"):
            continue
        par = Paragraph(filho, doc)
        if par.style.name == ESTILO_LISTA:
            continue               # LISTA DE GRAFICOS, no inicio do documento
        m = _RE_LEGENDA.match(par.text.strip())
        if not m:
            continue
        numero = int(m.group(1))
        if numero in vistos:
            continue
        vistos[numero] = _classificar(numero, m.group(2).strip())

    return tuple(vistos[n] for n in sorted(vistos))


def por_tipo(tipo: str, caminho_template: str = TEMPLATE) -> list:
    return [g for g in catalogo(caminho_template) if g.tipo == tipo]


def geraveis(caminho_template: str = TEMPLATE) -> list:
    """Os graficos que as duas planilhas da referencia sustentam."""
    return [g for g in catalogo(caminho_template) if g.geravel]


def pendentes(caminho_template: str = TEMPLATE) -> list:
    """Os graficos que dependem de dado que ainda nao chegou."""
    return [g for g in catalogo(caminho_template) if not g.geravel]


# ──────────────────────────────────────────────────────────────
# TABELAS E BLOCOS DE TEXTO (gerados por gerar_template_docx.py)
# ──────────────────────────────────────────────────────────────
_RE_TABELA = re.compile(r"^Tabela\s+(\d+)\s*[-–]\s*(.*)$")
PASTA_TEXTOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "materiais", "textos")


@lru_cache(maxsize=1)
def tabelas(caminho_template: str = TEMPLATE) -> tuple:
    """(numero, titulo) de cada tabela do template, pela legenda, sem repetir.

    Nenhuma tabela e preenchida pelo app ainda: todas mantem os valores da
    referencia de origem do template ate a planilha das tabelas ser mapeada.
    """
    from docx import Document

    vistos, lista = set(), []
    for paragrafo in Document(caminho_template).paragraphs:
        if paragrafo.style.name not in ESTILO_LEGENDA:
            continue
        m = _RE_TABELA.match(paragrafo.text.strip())
        if not m:
            continue
        numero = int(m.group(1))
        if numero in vistos:
            continue
        vistos.add(numero)
        lista.append((numero, re.sub(r"\s*\(\d+/\d+\)$", "", m.group(2).strip())))
    return tuple(lista)


@lru_cache(maxsize=4)
def blocos_texto(referencia: str = None) -> dict:
    """Catalogo dos blocos editaveis e seus textos na referencia de origem.

    Sem `referencia`, usa o arquivo mais recente de materiais/textos. Devolve
    o json como dicionario ({} quando nao ha catalogo).
    """
    import json

    if not os.path.isdir(PASTA_TEXTOS):
        return {}
    arquivos = sorted(a for a in os.listdir(PASTA_TEXTOS) if a.startswith("blocos_") and a.endswith(".json"))
    if referencia:
        arquivos = [a for a in arquivos if a == f"blocos_{referencia}.json"]
    if not arquivos:
        return {}
    with open(os.path.join(PASTA_TEXTOS, arquivos[-1]), encoding="utf-8") as f:
        return json.load(f)
