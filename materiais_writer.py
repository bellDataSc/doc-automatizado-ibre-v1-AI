"""Gera o relatorio Word do Tomo 1 - Materiais a partir do template do DNIT.

O template TEMPLATE_DNIT_Tomo-1_Materiais-SICRO.docx e gerado por
gerar_template_docx.py a partir do Word revisado da referencia anterior. Ele
traz a capa, todos os textos estaticos, as tabelas e as legendas dos 63
graficos na ordem da entrega. Este modulo:

  1. preenche os marcadores {{...}} da capa, dos cabecalhos/rodapes e do
     corpo (referencia atual e anterior, ordinal, numero do relatorio);
  2. escreve, em cada bloco de texto editavel (bookmark TXT_<id>_P<n>), o
     texto decidido pelo analista na tela;
  3. coloca cada grafico conferido no paragrafo de figura que vem logo
     depois da legenda correspondente;
  4. preenche as tabelas que a base do Snowflake sustenta (materiais_tabelas),
     mantendo a formatacao das linhas do template.

O casamento dos graficos e das tabelas e pela LEGENDA ("Grafico 12 - ...",
"Tabela 10 - ..."), nao pela posicao no documento: assim um grafico que nao
pode ser gerado nao desloca os demais. As tabelas que a base nao sustenta
ficam com os valores do template.
"""

from __future__ import annotations

import copy
import os
import re
import shutil
from dataclasses import dataclass, field

from docx.oxml.ns import qn
from docx.shared import Inches

from materiais_catalogo import ESTILO_FIGURA, ESTILO_LEGENDA, ESTILO_LISTA, TEMPLATE
from materiais_charts import figura_para_png

_RE_LEGENDA = re.compile(r"^Gráfico\s+(\d+)\s*[-–]\s*")

# Marcadores do template. A chave e o nome dentro de {{ }}.
MARCADORES_CAPA = (
    "NUM_RELATORIO",
    "REFERENCIA",
    "PROTOCOLO_FGV",
    "MES_ANO_EMISSAO",
    "ANO",
)
MARCADORES_CORPO = (
    "REF_ATUAL_EXTENSO",      # "outubro de 2026"
    "REF_ANTERIOR_EXTENSO",   # "julho de 2026"
    "REF_ATUAL_ARQUIVO",      # "Out.2026", como nos nomes dos arquivos
    "ORDINAL",                # "2" em "2ºRR", "2ºOrigem-Precos"
)
MARCADORES = MARCADORES_CAPA + MARCADORES_CORPO
_RE_MARCADOR = re.compile(r"\{\{([A-Z_]+)\}\}")


@dataclass
class ResultadoRelatorio:
    caminho: str
    graficos_inseridos: list = field(default_factory=list)
    graficos_ausentes: list = field(default_factory=list)
    marcadores_preenchidos: list = field(default_factory=list)
    marcadores_restantes: list = field(default_factory=list)
    blocos_escritos: list = field(default_factory=list)
    blocos_sem_bookmark: list = field(default_factory=list)
    tabelas_preenchidas: list = field(default_factory=list)
    tabelas_sem_slot: list = field(default_factory=list)
    tabelas_avisos: list = field(default_factory=list)


# ──────────────────────────────────────────────────────────────
# MARCADORES {{...}}
# ──────────────────────────────────────────────────────────────
def _texto_do_paragrafo(paragrafo) -> str:
    return "".join(t.text or "" for t in paragrafo.iter(qn("w:t")))


def _escrever_no_paragrafo(p, texto: str) -> None:
    """Escreve o texto no primeiro w:t e esvazia os demais.

    Usado quando o marcador esta quebrado em varios runs; o estilo do
    paragrafo sobrevive, a formatacao de caractere dos demais runs nao.
    """
    ts = list(p.iter(qn("w:t")))
    if not ts:
        return
    ts[0].text = texto
    ts[0].set(qn("xml:space"), "preserve")
    for t in ts[1:]:
        t.text = ""


def _substituir_no_paragrafo(p, valores: dict, preenchidos: list, restantes: list) -> None:
    texto = _texto_do_paragrafo(p)
    if "{{" not in texto:
        return

    # 1. run a run, preservando a formatacao de cada run
    for t in p.iter(qn("w:t")):
        if not t.text or "{{" not in t.text:
            continue

        def troca(m):
            chave = m.group(1)
            valor = valores.get(chave)
            if valor:
                preenchidos.append(chave)
                return str(valor)
            restantes.append(chave)
            return m.group(0)

        t.text = _RE_MARCADOR.sub(troca, t.text)

    # 2. marcador partido entre runs: reescreve o paragrafo inteiro
    texto = _texto_do_paragrafo(p)
    if "{{" in texto:
        sobra = set(_RE_MARCADOR.findall(texto))
        novo = texto
        for chave in sobra:
            valor = valores.get(chave)
            if valor:
                novo = novo.replace("{{" + chave + "}}", str(valor))
                preenchidos.append(chave)
            else:
                restantes.append(chave)
        if novo != texto:
            _escrever_no_paragrafo(p, novo)


def _partes_de_cabecalho(documento):
    vistos = set()
    for secao in documento.sections:
        for hf in (secao.header, secao.footer, secao.first_page_header,
                   secao.first_page_footer, secao.even_page_header,
                   secao.even_page_footer):
            parte = hf.part
            if id(parte) in vistos:
                continue
            vistos.add(id(parte))
            yield parte


def preencher_marcadores(documento, valores: dict) -> tuple:
    """Troca {{MARCADOR}} pelo valor em corpo, tabelas, cabecalhos e rodapes.

    Devolve (preenchidos, restantes) com as chaves encontradas.
    """
    preenchidos, restantes = [], []
    for p in documento.element.body.iter(qn("w:p")):
        _substituir_no_paragrafo(p, valores, preenchidos, restantes)
    for parte in _partes_de_cabecalho(documento):
        for p in parte.element.iter(qn("w:p")):
            _substituir_no_paragrafo(p, valores, preenchidos, restantes)
    return preenchidos, restantes


# compatibilidade com o nome antigo
preencher_capa = preencher_marcadores


# ──────────────────────────────────────────────────────────────
# BLOCOS DE TEXTO EDITAVEIS (bookmarks TXT_<id>_P<n>)
# ──────────────────────────────────────────────────────────────
_RE_BOOKMARK = re.compile(r"^TXT_(B\d{3})_P(\d+)$")


def _paragrafos_dos_blocos(documento) -> dict:
    """{id do bloco: [paragrafo xml, ...]} na ordem do documento."""
    blocos: dict = {}
    for p in documento.element.body.iter(qn("w:p")):
        for marca in p.findall(qn("w:bookmarkStart")):
            m = _RE_BOOKMARK.match(marca.get(qn("w:name"), ""))
            if m:
                blocos.setdefault(m.group(1), []).append(p)
    return blocos


def _escrever_texto_no_paragrafo(p, texto: str) -> None:
    """Mantem o primeiro run de texto (com a formatacao) e remove os demais."""
    runs = [r for r in p.findall(qn("w:r")) if r.findall(qn("w:t"))]
    if not runs:
        r = p.makeelement(qn("w:r"), {})
        fim = p.find(qn("w:bookmarkEnd"))
        if fim is not None:
            fim.addprevious(r)
        else:
            p.append(r)
        t = r.makeelement(qn("w:t"), {})
        r.append(t)
        runs = [r]
    for r in runs[1:]:
        p.remove(r)
    ts = runs[0].findall(qn("w:t"))
    ts[0].text = texto
    ts[0].set(qn("xml:space"), "preserve")
    for t in ts[1:]:
        runs[0].remove(t)


def preencher_blocos(documento, textos: dict) -> tuple:
    """Escreve o texto de cada bloco nos paragrafos marcados pelo bookmark.

    `textos` mapeia o id do bloco (B001...) para o texto final; quebras de
    linha separam paragrafos. Sobrando paragrafos no template, sao
    removidos; faltando, o ultimo e duplicado. Devolve (escritos, sem_bookmark).
    """
    por_bloco = _paragrafos_dos_blocos(documento)
    escritos, sem_bookmark = [], []

    for bloco_id, texto in textos.items():
        paragrafos = por_bloco.get(bloco_id)
        if not paragrafos:
            sem_bookmark.append(bloco_id)
            continue
        linhas = [l.strip() for l in (texto or "").split("\n") if l.strip()] or [""]

        while len(paragrafos) < len(linhas):
            clone = copy.deepcopy(paragrafos[-1])
            paragrafos[-1].addnext(clone)
            paragrafos.append(clone)
        for p in paragrafos[len(linhas):]:
            p.getparent().remove(p)
        for p, linha in zip(paragrafos, linhas):
            _escrever_texto_no_paragrafo(p, linha)
        escritos.append(bloco_id)

    return escritos, sem_bookmark


# ──────────────────────────────────────────────────────────────
# FIGURAS
# ──────────────────────────────────────────────────────────────
def _paragrafos_do_corpo(documento):
    from docx.text.paragraph import Paragraph

    for filho in documento.element.body.iterchildren():
        if filho.tag == qn("w:p"):
            yield Paragraph(filho, documento)


def _largura_por_paragrafo(documento) -> dict:
    """Largura util da secao de cada paragrafo do corpo, em polegadas.

    A quebra de secao mora no w:pPr/w:sectPr do ultimo paragrafo da secao;
    a ultima secao mora no w:sectPr do corpo. Importa porque o relatorio
    alterna paginas em retrato e em paisagem.
    """
    secoes = documento.sections
    larguras = []
    for secao in secoes:
        try:
            larguras.append(
                (secao.page_width - secao.left_margin - secao.right_margin) / 914400
            )
        except TypeError:
            larguras.append(6.3)

    mapa, indice = {}, 0
    for paragrafo in _paragrafos_do_corpo(documento):
        mapa[id(paragrafo._p)] = larguras[min(indice, len(larguras) - 1)]
        propriedades = paragrafo._p.find(qn("w:pPr"))
        if propriedades is not None and propriedades.find(qn("w:sectPr")) is not None:
            indice += 1
    return mapa


def _slot_da_figura(paragrafos: list, inicio: int):
    """Paragrafo de figura vazio logo depois da legenda, ou None.

    Duas legendas do template (Graficos 12 e 47) nao tem paragrafo de
    figura; nesse caso quem chama cria um.
    """
    for indice in range(inicio + 1, min(inicio + 4, len(paragrafos))):
        paragrafo = paragrafos[indice]
        if _RE_LEGENDA.match(paragrafo.text.strip()):
            return None
        if paragrafo.style.name == ESTILO_FIGURA and not paragrafo.text.strip():
            return paragrafo
    return None


def _novo_paragrafo_depois(paragrafo, documento):
    from docx.text.paragraph import Paragraph

    elemento = paragrafo._p.makeelement(qn("w:p"), {})
    paragrafo._p.addnext(elemento)
    novo = Paragraph(elemento, documento)
    try:
        novo.style = documento.styles[ESTILO_FIGURA]
    except KeyError:
        pass
    return novo


def _imagem(entrada, dpi: int):
    """Aceita uma figura matplotlib ou um PNG ja pronto."""
    if hasattr(entrada, "savefig"):
        return figura_para_png(entrada, dpi)
    entrada.seek(0)
    return entrada


def inserir_graficos(documento, figuras_por_numero: dict, dpi: int = 200) -> tuple:
    """Coloca cada grafico no slot que segue a sua legenda.

    `figuras_por_numero` mapeia o numero do grafico para uma figura
    matplotlib ou para o PNG ja gerado na conferencia. Devolve
    (inseridos, ausentes).
    """
    larguras = _largura_por_paragrafo(documento)
    paragrafos = [
        p for p in _paragrafos_do_corpo(documento) if p.style.name != ESTILO_LISTA
    ]

    inseridos, vistos = [], set()

    for indice, paragrafo in enumerate(paragrafos):
        if paragrafo.style.name not in ESTILO_LEGENDA:
            continue
        casamento = _RE_LEGENDA.match(paragrafo.text.strip())
        if not casamento:
            continue
        numero = int(casamento.group(1))
        if numero in vistos:
            continue
        vistos.add(numero)

        figura = figuras_por_numero.get(numero)
        if figura is None:
            continue

        slot = _slot_da_figura(paragrafos, indice)
        if slot is None:
            slot = _novo_paragrafo_depois(paragrafo, documento)
            largura = larguras.get(id(paragrafo._p), 6.3)
        else:
            largura = larguras.get(id(slot._p), 6.3)

        for run in slot.runs:
            run.text = ""
        slot.add_run().add_picture(_imagem(figura, dpi), width=Inches(largura))
        inseridos.append(numero)

    ausentes = sorted(set(figuras_por_numero) - set(inseridos))
    return sorted(inseridos), ausentes


# ──────────────────────────────────────────────────────────────
# TABELAS
# ──────────────────────────────────────────────────────────────
_RE_LEGENDA_TABELA = re.compile(r"^Tabela\s+(\d+)\s*[-–]")


def _celulas_unicas(linha):
    """Celulas da linha sem repetir as mescladas horizontalmente."""
    saida, anterior = [], None
    for celula in linha.cells:
        if celula._tc is anterior:
            continue
        anterior = celula._tc
        saida.append(celula)
    return saida


def _texto_celula(celula) -> str:
    return " ".join(p.text for p in celula.paragraphs).strip()


def _escrever_celula(celula, texto: str) -> None:
    """Mantem o primeiro paragrafo e o primeiro run (formatacao); apaga o resto."""
    paragrafos = celula._tc.findall(qn("w:p"))
    for p in paragrafos[1:]:
        celula._tc.remove(p)
    _escrever_texto_no_paragrafo(paragrafos[0], texto)


def tabelas_do_template(documento) -> dict:
    """{numero: [Table, ...]} na ordem do documento, pela legenda.

    A legenda normalmente vem antes da tabela. A excecao e a tabela 15, cuja
    legenda vem depois: uma legenda sem tabela a seguir reivindica a tabela
    imediatamente anterior quando esta ja tem outra dona.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    sequencia = []        # ("leg", numero) ou ("tbl", Table)
    for filho in documento.element.body.iterchildren():
        if filho.tag == qn("w:p"):
            par = Paragraph(filho, documento)
            if par.style.name == ESTILO_LISTA:
                continue
            m = _RE_LEGENDA_TABELA.match(par.text.strip())
            if m:
                sequencia.append(("leg", int(m.group(1)), filho))
        elif filho.tag == qn("w:tbl"):
            sequencia.append(("tbl", Table(filho, documento), filho))

    donos = []            # numero de cada tabela, na ordem
    atual = None
    for i, (tipo, valor, _el) in enumerate(sequencia):
        if tipo == "leg":
            atual = valor
            seguinte = next((t for t, _v, _e in sequencia[i + 1:i + 2]), None)
            if seguinte != "tbl" and i > 0 and sequencia[i - 1][0] == "tbl":
                anterior = sequencia[i - 1][1]
                if donos and donos[-1][0] != valor and donos[-1][0] in [d[0] for d in donos[:-1]]:
                    donos[-1] = (valor, anterior, sequencia[i - 1][2])
        else:
            donos.append((atual, valor, _el))

    saida: dict = {}
    for numero, tabela, _el in donos:
        if numero is not None:
            saida.setdefault(numero, []).append(tabela)
    return saida


def linhas_da_tabela(tabela) -> list:
    return [[_texto_celula(c) for c in _celulas_unicas(l)] for l in tabela.rows]


def _remover_parte(tabela) -> None:
    """Remove uma parte vazia: a tabela, a legenda antes e o 'Fonte' depois."""
    from docx.text.paragraph import Paragraph

    elemento = tabela._tbl
    anterior = elemento.getprevious()
    proximo = elemento.getnext()
    pai = elemento.getparent()
    if anterior is not None and anterior.tag == qn("w:p") and \
            _RE_LEGENDA_TABELA.match(Paragraph(anterior, None).text.strip()):
        pai.remove(anterior)
    if proximo is not None and proximo.tag == qn("w:p") and \
            Paragraph(proximo, None).text.strip().lower().startswith("fonte"):
        pai.remove(proximo)
    pai.remove(elemento)


def _ajustar_linhas(tabela, inicio: int, fim: int, quantidade: int) -> list:
    """Deixa exatamente `quantidade` linhas entre inicio e fim (exclusivo).

    Clona a primeira linha do intervalo para crescer; remove do fim para
    encolher. Devolve as linhas (python-docx Row) do intervalo.
    """
    trs = [l._tr for l in tabela.rows]
    modelo = copy.deepcopy(trs[inicio])
    existentes = trs[inicio:fim]
    while len(existentes) < quantidade:
        novo = copy.deepcopy(modelo)
        existentes[-1].addnext(novo)
        existentes.append(novo)
    for tr in existentes[quantidade:]:
        tr.getparent().remove(tr)
    return list(tabela.rows)[inicio:inicio + quantidade]


def _limpar_sombreamento(celula) -> None:
    """Tira o fundo da celula. Nas tabelas item x UF do template as celulas
    vem pintadas (vermelho/azul) sem regra consistente; como o valor muda,
    a cor antiga nao pode ficar."""
    tcpr = celula._tc.find(qn("w:tcPr"))
    if tcpr is None:
        return
    for shd in tcpr.findall(qn("w:shd")):
        tcpr.remove(shd)


def _preencher_linhas(partes: list, linhas: list, cabecalho: int = 1,
                      limpar_fundo_desde: int = None) -> None:
    """Distribui `linhas` pelas partes da tabela, parte a parte, mantendo a
    capacidade de cada uma; a ultima parte cresce ou encolhe; partes que
    ficam vazias sao removidas."""
    restantes = list(linhas)
    for indice, tabela in enumerate(partes):
        ultima = indice == len(partes) - 1
        capacidade = len(tabela.rows) - cabecalho
        lote = restantes if ultima else restantes[:capacidade]
        restantes = restantes[len(lote):]
        if not lote:
            _remover_parte(tabela)
            continue
        for linha, valores in zip(_ajustar_linhas(tabela, cabecalho, len(tabela.rows), len(lote)), lote):
            for posicao, (celula, valor) in enumerate(zip(_celulas_unicas(linha), valores)):
                _escrever_celula(celula, valor)
                if limpar_fundo_desde is not None and posicao >= limpar_fundo_desde:
                    _limpar_sombreamento(celula)


def _preencher_faixas(partes: list, pronta) -> None:
    """Tabelas de distribuicao: os rotulos das faixas sao os do Word; a base
    traz quantidade e proporcao na ORDEM do Word. Os totais seguem o que o
    rotulo da linha diz."""
    from materiais_tabelas import _RE_INSUMOS, _RE_TOTAL, _RE_TOTAL_GERAL, inteiro

    slots, totais = [], []
    for tabela in partes:
        for linha in list(tabela.rows)[1:]:
            celulas = _celulas_unicas(linha)
            rotulo = _texto_celula(celulas[0])
            if _RE_TOTAL.match(rotulo) or _RE_INSUMOS.match(rotulo):
                totais.append((rotulo, celulas, len(slots)))
            else:
                slots.append(celulas)
    quantidades = pronta.rodape.get("quantidades", [])
    for celulas, valores in zip(slots, pronta.linhas):
        for celula, valor in zip(celulas[1:], valores):
            if _texto_celula(celula) == "" and celula is not celulas[1]:
                continue          # celula vazia no template (ex.: proporcao dos zerados)
            _escrever_celula(celula, valor)
    if len(pronta.linhas) != len(slots):
        pronta.avisos.append(f"a base trouxe {len(pronta.linhas)} faixa(s) e o Word tem {len(slots)}")
    ultimo_total = None
    for rotulo, celulas, posicao in totais:
        if _RE_TOTAL_GERAL.match(rotulo):
            valor = sum(quantidades)
        elif _RE_INSUMOS.match(rotulo):
            valor = (ultimo_total / 27) if ultimo_total is not None else None
        else:
            valor = sum(quantidades[:posicao])
            ultimo_total = valor
        if valor is not None and len(celulas) > 1:
            _escrever_celula(celulas[1], inteiro(valor))


def _preencher_historico(tabela, pronta) -> None:
    """Tabela 26: cabecalho com as referencias, pares quantidade/proporcao
    por faixa e a linha de total."""
    linhas = list(tabela.rows)
    cabecalho = _celulas_unicas(linhas[1])
    for celula, rotulo in zip(cabecalho[1:], pronta.cabecalho):
        _escrever_celula(celula, rotulo)
    dados = [l for l in linhas[2:] if not _texto_celula(_celulas_unicas(l)[0]).lower().startswith("total")]
    for linha, valores in zip(dados, pronta.linhas):
        celulas = _celulas_unicas(linha)
        for celula, valor in zip(celulas[2:], valores):
            _escrever_celula(celula, valor)
    if len(dados) != len(pronta.linhas):
        pronta.avisos.append(f"a base trouxe {len(pronta.linhas)} linha(s) e o Word tem {len(dados)}")
    total = [l for l in linhas[2:] if _texto_celula(_celulas_unicas(l)[0]).lower().startswith("total")]
    if total:
        celulas = _celulas_unicas(total[0])
        for celula, valor in zip(celulas[1:], pronta.rodape.get("total", [])):
            _escrever_celula(celula, valor)


def _preencher_fixa(tabela, pronta) -> None:
    linhas = list(tabela.rows)
    for (i, j), texto in pronta.celulas.items():
        if i < len(linhas):
            celulas = _celulas_unicas(linhas[i])
            if j < len(celulas) and texto:
                _escrever_celula(celulas[j], texto)


def preencher_tabelas(documento, prontas: dict) -> tuple:
    """Escreve cada TabelaPronta na(s) parte(s) da tabela do template.

    Devolve (preenchidas, sem_tabela, avisos) com os numeros.
    """
    por_numero = tabelas_do_template(documento)
    preenchidas, sem_tabela, avisos = [], [], []
    for numero, pronta in sorted(prontas.items()):
        partes = por_numero.get(numero)
        if not partes:
            sem_tabela.append(numero)
            continue
        if pronta.tipo == "erro":
            avisos.extend(f"Tabela {numero}: {a}" for a in pronta.avisos)
            continue
        if pronta.tipo in ("lista", "resumo"):
            _preencher_linhas(partes, pronta.linhas)
        elif pronta.tipo == "uf":
            _preencher_linhas(partes, pronta.linhas, limpar_fundo_desde=2)
        elif pronta.tipo == "faixas":
            _preencher_faixas(partes, pronta)
        elif pronta.tipo == "historico":
            _preencher_historico(partes[0], pronta)
        elif pronta.tipo == "fixa":
            _preencher_fixa(partes[0], pronta)
        else:
            continue
        preenchidas.append(numero)
        avisos.extend(f"Tabela {numero}: {a}" for a in pronta.avisos)
    return preenchidas, sem_tabela, avisos


# ──────────────────────────────────────────────────────────────
# GERACAO
# ──────────────────────────────────────────────────────────────
def gerar_relatorio(
    figuras_por_numero: dict,
    caminho_saida: str,
    capa: dict = None,
    caminho_template: str = TEMPLATE,
    dpi: int = 200,
    textos: dict = None,
    tabelas: dict = None,
) -> ResultadoRelatorio:
    """Copia o template, preenche marcadores, textos e tabelas e insere os graficos.

    `capa` traz os valores de todos os marcadores (capa e corpo);
    `textos` traz o texto final de cada bloco editavel, por id;
    `tabelas` traz as TabelaPronta (materiais_tabelas.montar), por numero.
    """
    from docx import Document

    if not os.path.exists(caminho_template):
        raise FileNotFoundError(f"template não encontrado: {caminho_template}")

    pasta = os.path.dirname(caminho_saida)
    if pasta:
        os.makedirs(pasta, exist_ok=True)
    shutil.copyfile(caminho_template, caminho_saida)

    documento = Document(caminho_saida)
    escritos, sem_bookmark = preencher_blocos(documento, textos or {})
    preenchidos, restantes = preencher_marcadores(documento, capa or {})
    tab_ok, tab_sem, tab_avisos = preencher_tabelas(documento, tabelas or {})
    inseridos, ausentes = inserir_graficos(documento, figuras_por_numero, dpi)
    documento.save(caminho_saida)

    return ResultadoRelatorio(
        caminho=caminho_saida,
        graficos_inseridos=inseridos,
        graficos_ausentes=ausentes,
        marcadores_preenchidos=sorted(set(preenchidos)),
        marcadores_restantes=sorted(set(restantes)),
        blocos_escritos=escritos,
        blocos_sem_bookmark=sem_bookmark,
        tabelas_preenchidas=tab_ok,
        tabelas_sem_slot=tab_sem,
        tabelas_avisos=tab_avisos,
    )
